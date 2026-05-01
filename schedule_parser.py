import json
import os
import anthropic
from docx import Document


def extract_text_from_docx(path: str) -> str:
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = ' | '.join(cell.text.strip() for cell in row.cells)
            if row_text.strip(' |'):
                parts.append(row_text)
    return '\n'.join(parts)


SHIFT_TYPE_DESCRIPTIONS = """
סוגי משמרות אפשריים:
- morning_weekday: בוקר חול (ביום חול רגיל)
- morning_shabbat: בוקר שבת (בשבת או חג)
- afternoon_weekday: צהריים חול (ביום חול רגיל)
- afternoon_friday: צהריים שישי (ביום שישי לפני שבת)
- afternoon_shabbat: צהריים שבת (בשבת או חג)
- night_weekday: לילה חול (לילה ביום חול)
- night_friday_shabbat: לילה שישי/שבת (לילה שישי-שבת או שבת-ראשון)
- double_morning_afternoon_weekday: כפולה בוקר+צהריים חול
- double_morning_afternoon_shabbat: כפולה בוקר+צהריים שבת
- double_afternoon_night_weekday: כפולה צהריים+לילה חול
- double_afternoon_night_friday_shabbat: כפולה צהריים+לילה שישי/שבת
"""

RESPONSE_SCHEMA = {
    "employee_name": "שם העובד (ריק אם לא נמצא)",
    "month": "חודש ושנה (למשל: אפריל 2026)",
    "shift_counts": {
        "morning_weekday": 0,
        "morning_shabbat": 0,
        "afternoon_weekday": 0,
        "afternoon_friday": 0,
        "afternoon_shabbat": 0,
        "night_weekday": 0,
        "night_friday_shabbat": 0,
        "double_morning_afternoon_weekday": 0,
        "double_morning_afternoon_shabbat": 0,
        "double_afternoon_night_weekday": 0,
        "double_afternoon_night_friday_shabbat": 0,
    },
    "total_shifts": 0,
    "notes": "הערות חשובות",
}


def parse_schedule_with_claude(text: str, employee_name: str = '') -> dict:
    client = anthropic.Anthropic()

    name_instruction = f'חפש ספציפית את העובד בשם "{employee_name}" בסידור.' if employee_name else \
        'אם יש מספר עובדים בסידור, ספור את המשמרות של כולם ביחד אלא אם ברור שמדובר בעובד ספציפי.'

    prompt = f"""אתה מנתח סידורי עבודה לבקרי מקורות. {name_instruction}

{SHIFT_TYPE_DESCRIPTIONS}

תוכן הסידור:
---
{text}
---

ספור את מספר המשמרות מכל סוג שהעובד/ים יבצעו בחודש זה.
שים לב:
- "כפולה" = משמרת ארוכה שמכסה שתי משמרות
- שישי = יום שישי (ערב שבת)
- שבת = שבת (כולל חגים)
- חול = ימים א׳-ה׳

החזר JSON בלבד, ללא הסברים, בפורמט זה:
{json.dumps(RESPONSE_SCHEMA, ensure_ascii=False, indent=2)}"""

    message = client.messages.create(
        model='claude-opus-4-7',
        max_tokens=1024,
        messages=[{'role': 'user', 'content': prompt}],
    )

    text_response = message.content[0].text.strip()
    if text_response.startswith('```'):
        lines = text_response.split('\n')
        text_response = '\n'.join(lines[1:-1])

    data = json.loads(text_response)
    total = sum(data.get('shift_counts', {}).values())
    data['total_shifts'] = total
    return data


def parse_schedule(file_path: str, employee_name: str = '') -> dict:
    text = extract_text_from_docx(file_path)
    if not text.strip():
        raise ValueError('לא נמצא תוכן בקובץ Word')
    return parse_schedule_with_claude(text, employee_name)
